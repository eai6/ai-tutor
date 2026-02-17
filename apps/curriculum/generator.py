"""
Curriculum Auto-Generator

This module processes uploaded curriculum documents (PDF, DOCX) and automatically:
1. Extracts topics and learning objectives
2. Generates lessons with proper structure
3. Creates practice questions and exit tickets
4. Generates educational media (images via DALL-E)

Usage:
    python manage.py generate_curriculum --file syllabus.pdf --subject Geography
    
Or via API:
    POST /api/curriculum/generate/
    {
        "file_id": 123,
        "subject": "Geography",
        "grade_level": "S1-S3"
    }
"""

import os
import json
import re
import time
from typing import List, Dict, Optional, Tuple
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.core.files.base import ContentFile

from apps.curriculum.models import Course, Unit, Lesson, LessonStep
from apps.accounts.models import Institution
from apps.media_library.models import MediaAsset, StepMedia


class CurriculumGenerator:
    """
    Generates complete curriculum from syllabus documents.
    
    Flow:
    1. Parse document → Extract raw text
    2. Analyze with Claude → Identify units, topics, objectives
    3. Generate lessons → Create structured content for each topic
    4. Generate questions → Practice problems and exit tickets
    5. Generate media → Educational images with DALL-E
    """
    
    # Seychelles-specific context for content generation
    LOCAL_CONTEXT = """
    Use Seychelles context where appropriate:
    - Location: Victoria (capital), Mahé, Praslin, La Digue islands
    - Currency: Seychelles Rupee (SCR)
    - Local names: Jean, Marie, Pierre, Ansel, Lisette, Michel
    - Industries: Tourism, fishing, cinnamon, coconut, vanilla
    - Geography: Granite islands, coral atolls, tropical climate
    - Culture: Creole heritage, multicultural (African, Asian, European)
    """
    
    def __init__(self, institution: Institution, anthropic_key: str, openai_key: str = None):
        self.institution = institution
        self.anthropic_key = anthropic_key
        self.openai_key = openai_key
        
        # Initialize clients
        import anthropic
        self.claude = anthropic.Anthropic(api_key=anthropic_key)
        
        if openai_key:
            from openai import OpenAI
            self.dalle = OpenAI(api_key=openai_key)
        else:
            self.dalle = None
    
    def generate_from_document(
        self,
        document_path: str,
        subject: str,
        grade_level: str = "S1-S5",
        generate_media: bool = True
    ) -> Course:
        """
        Main entry point: Generate full curriculum from a document.
        
        Args:
            document_path: Path to syllabus PDF or DOCX
            subject: Subject name (e.g., "Geography")
            grade_level: Grade level (e.g., "S1-S5")
            generate_media: Whether to generate images with DALL-E
            
        Returns:
            Created Course object with all units, lessons, and content
        """
        
        print(f"📚 Generating curriculum for {subject} from {document_path}")
        
        # Step 1: Extract text from document
        print("  1. Extracting text from document...")
        document_text = self._extract_document_text(document_path)
        
        # Step 2: Analyze structure with Claude
        print("  2. Analyzing curriculum structure...")
        structure = self._analyze_curriculum_structure(document_text, subject)
        
        # Step 3: Create course and units
        print("  3. Creating course structure...")
        course = self._create_course_structure(subject, grade_level, structure)
        
        # Step 4: Generate lesson content
        print("  4. Generating lesson content...")
        self._generate_lesson_content(course, structure)
        
        # Step 5: Generate media (if enabled)
        if generate_media and self.dalle:
            print("  5. Generating educational media...")
            self._generate_lesson_media(course)
        
        print(f"✅ Curriculum generation complete!")
        print(f"   Created: {course.units.count()} units, {Lesson.objects.filter(unit__course=course).count()} lessons")
        
        return course
    
    def _extract_document_text(self, document_path: str) -> str:
        """Extract text from PDF or DOCX."""
        
        path = Path(document_path)
        
        if path.suffix.lower() == '.pdf':
            return self._extract_pdf_text(document_path)
        elif path.suffix.lower() in ['.docx', '.doc']:
            return self._extract_docx_text(document_path)
        elif path.suffix.lower() == '.txt':
            with open(document_path, 'r', encoding='utf-8') as f:
                return f.read()
        else:
            raise ValueError(f"Unsupported file type: {path.suffix}")
    
    def _extract_pdf_text(self, path: str) -> str:
        """Extract text from PDF using pypdf."""
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text
        except ImportError:
            # Fallback to basic extraction
            import subprocess
            result = subprocess.run(
                ['pdftotext', '-layout', path, '-'],
                capture_output=True, text=True
            )
            return result.stdout
    
    def _extract_docx_text(self, path: str) -> str:
        """Extract text from DOCX."""
        try:
            from docx import Document
            doc = Document(path)
            return "\n".join([para.text for para in doc.paragraphs])
        except ImportError:
            raise ImportError("python-docx not installed. Run: pip install python-docx")
    
    def _analyze_curriculum_structure(self, document_text: str, subject: str) -> Dict:
        """Use Claude to analyze curriculum and extract structure."""
        
        prompt = f"""Analyze this curriculum document for {subject} and extract the structure.

DOCUMENT:
{document_text[:15000]}  # Limit to avoid token limits

Extract and return as JSON:
{{
    "subject": "{subject}",
    "description": "Brief description of the curriculum",
    "units": [
        {{
            "title": "Unit title",
            "description": "Unit description",
            "topics": [
                {{
                    "title": "Topic/Lesson title",
                    "objective": "Learning objective",
                    "key_concepts": ["concept1", "concept2"],
                    "skills": ["skill1", "skill2"]
                }}
            ]
        }}
    ]
}}

Guidelines:
- Extract ALL units and topics mentioned
- Each topic should have a clear learning objective
- Identify 3-5 key concepts per topic
- Group related topics into units logically
- Use exact terminology from the document where possible

Return ONLY valid JSON, no other text."""

        response = self.claude.messages.create(
            model="claude-sonnet-4-20250514",
            max_tokens=4096,
            messages=[{"role": "user", "content": prompt}]
        )
        
        content = response.content[0].text.strip()
        
        # Clean up JSON
        if content.startswith('```'):
            content = content.split('```')[1]
            if content.startswith('json'):
                content = content[4:]
            content = content.strip()
        
        return json.loads(content)
    
    @transaction.atomic
    def _create_course_structure(self, subject: str, grade_level: str, structure: Dict) -> Course:
        """Create Course, Units from analyzed structure."""
        
        # Create or update course
        course, _ = Course.objects.update_or_create(
            institution=self.institution,
            title=subject,
            defaults={
                'description': structure.get('description', f'{subject} curriculum'),
                'grade_level': grade_level,
                'is_published': True,
            }
        )
        
        # Create units
        for i, unit_data in enumerate(structure.get('units', [])):
            unit, _ = Unit.objects.update_or_create(
                course=course,
                title=unit_data['title'],
                defaults={
                    'description': unit_data.get('description', ''),
                    'order_index': i,
                }
            )
            
            # Create lesson placeholders
            for j, topic in enumerate(unit_data.get('topics', [])):
                Lesson.objects.update_or_create(
                    unit=unit,
                    title=topic['title'],
                    defaults={
                        'objective': topic['objective'],
                        'estimated_minutes': 20,
                        'mastery_rule': Lesson.MasteryRule.PASS_QUIZ,
                        'order_index': j,
                        'is_published': True,
                        # Store extra data for content generation
                        'metadata': {
                            'key_concepts': topic.get('key_concepts', []),
                            'skills': topic.get('skills', []),
                        }
                    }
                )
        
        return course
    
    def _generate_lesson_content(self, course: Course, structure: Dict):
        """Generate detailed content for each lesson."""
        
        lessons = Lesson.objects.filter(unit__course=course).select_related('unit')
        
        for lesson in lessons:
            print(f"    Generating content for: {lesson.title}")
            
            # Find topic data from structure
            topic_data = self._find_topic_data(structure, lesson.title)
            
            # Generate comprehensive lesson content
            content = self._generate_single_lesson_content(lesson, topic_data)
            
            # Create lesson steps
            self._create_lesson_steps(lesson, content)
            
            # Rate limiting
            time.sleep(1)
    
    def _find_topic_data(self, structure: Dict, lesson_title: str) -> Dict:
        """Find topic data in structure by lesson title."""
        for unit in structure.get('units', []):
            for topic in unit.get('topics', []):
                if topic['title'] == lesson_title:
                    return topic
        return {}
    
    def _generate_single_lesson_content(self, lesson: Lesson, topic_data: Dict) -> Dict:
        """Generate all content for a single lesson."""
        
        prompt = f"""Generate comprehensive tutoring content for this lesson.

LESSON: {lesson.title}
UNIT: {lesson.unit.title}
OBJECTIVE: {lesson.objective}
KEY CONCEPTS: {topic_data.get('key_concepts', [])}
SKILLS: {topic_data.get('skills', [])}

{self.LOCAL_CONTEXT}

Generate content in this JSON structure:
{{
    "teaching_script": "Detailed explanation of the topic (3-4 paragraphs). Include:
        - Clear introduction
        - Main concepts explained simply
        - Real-world examples using Seychelles context
        - Key takeaways",
    
    "worked_example": {{
        "problem": "A concrete problem to solve",
        "solution_steps": ["Step 1...", "Step 2...", "Step 3..."],
        "final_answer": "The answer"
    }},
    
    "retrieval_questions": [
        {{
            "question": "Quick question about prerequisite knowledge",
            "options": ["A) ...", "B) ...", "C) ...", "D) ..."],
            "correct": "B",
            "explanation": "Why this is correct"
        }},
        {{
            "question": "Another retrieval question",
            "options": ["A) ...", "B) ...", "C) ...", "D) ..."],
            "correct": "A",
            "explanation": "Why this is correct"
        }}
    ],
    
    "practice_questions": [
        {{
            "question": "Practice problem 1 (easier)",
            "options": ["A) ...", "B) ...", "C) ...", "D) ..."] or null for free response,
            "correct": "C",
            "hints": ["Hint 1", "Hint 2", "Hint 3"],
            "explanation": "Full explanation"
        }},
        {{
            "question": "Practice problem 2 (medium)",
            "options": null,
            "correct": "Expected answer",
            "hints": ["Hint 1", "Hint 2"],
            "explanation": "Full explanation"
        }},
        {{
            "question": "Practice problem 3 (harder)",
            "options": ["A) ...", "B) ...", "C) ...", "D) ..."],
            "correct": "D",
            "hints": ["Hint 1", "Hint 2", "Hint 3"],
            "explanation": "Full explanation"
        }}
    ],
    
    "exit_ticket": [
        {{
            "question": "Assessment question 1",
            "options": ["A) ...", "B) ...", "C) ...", "D) ..."],
            "correct": "A",
            "explanation": "Brief explanation"
        }},
        // ... 5 questions total
    ],
    
    "summary": "2-3 sentence summary of key takeaways",
    
    "image_suggestions": [
        {{
            "title": "Descriptive title",
            "prompt": "Detailed DALL-E prompt for educational image",
            "purpose": "What this image helps explain"
        }}
    ]
}}

Make content engaging, age-appropriate for secondary students, and rooted in Seychelles context.
Return ONLY valid JSON."""

        response = self.claude.messages.create(
            model="claude-sonnet-4-20250514",
            max_tokens=4096,
            messages=[{"role": "user", "content": prompt}]
        )
        
        content = response.content[0].text.strip()
        
        # Clean up JSON
        if content.startswith('```'):
            content = content.split('```')[1]
            if content.startswith('json'):
                content = content[4:]
            content = content.strip()
        
        try:
            return json.loads(content)
        except json.JSONDecodeError as e:
            print(f"      Warning: Failed to parse content for {lesson.title}: {e}")
            return {}
    
    @transaction.atomic
    def _create_lesson_steps(self, lesson: Lesson, content: Dict):
        """Create LessonStep objects from generated content."""
        
        # Clear existing steps
        lesson.steps.all().delete()
        
        order_index = 0
        
        # Teaching content step
        if content.get('teaching_script'):
            LessonStep.objects.create(
                lesson=lesson,
                order_index=order_index,
                step_type=LessonStep.StepType.TEACH,
                teacher_script=content['teaching_script'],
                answer_type=LessonStep.AnswerType.NONE,
                metadata={
                    'retrieval_questions': content.get('retrieval_questions', []),
                    'worked_example': content.get('worked_example'),
                }
            )
            order_index += 1
        
        # Worked example step
        if content.get('worked_example'):
            example = content['worked_example']
            LessonStep.objects.create(
                lesson=lesson,
                order_index=order_index,
                step_type=LessonStep.StepType.WORKED_EXAMPLE,
                teacher_script=f"Problem: {example.get('problem', '')}\n\nSolution:\n" + 
                              "\n".join(example.get('solution_steps', [])) +
                              f"\n\nAnswer: {example.get('final_answer', '')}",
                answer_type=LessonStep.AnswerType.NONE,
            )
            order_index += 1
        
        # Practice questions
        for i, q in enumerate(content.get('practice_questions', [])[:3]):
            LessonStep.objects.create(
                lesson=lesson,
                order_index=order_index,
                step_type=LessonStep.StepType.PRACTICE,
                question=q['question'],
                expected_answer=q.get('correct', ''),
                choices=q.get('options'),
                hint_1=q.get('hints', [''])[0] if q.get('hints') else '',
                hint_2=q.get('hints', ['', ''])[1] if len(q.get('hints', [])) > 1 else '',
                hint_3=q.get('hints', ['', '', ''])[2] if len(q.get('hints', [])) > 2 else '',
                answer_type=LessonStep.AnswerType.MULTIPLE_CHOICE if q.get('options') else LessonStep.AnswerType.FREE_RESPONSE,
                max_attempts=3,
                metadata={'explanation': q.get('explanation', '')}
            )
            order_index += 1
        
        # Exit ticket questions
        for i, q in enumerate(content.get('exit_ticket', [])[:5]):
            LessonStep.objects.create(
                lesson=lesson,
                order_index=order_index,
                step_type=LessonStep.StepType.QUIZ,
                question=q['question'],
                expected_answer=q.get('correct', ''),
                choices=q.get('options'),
                answer_type=LessonStep.AnswerType.MULTIPLE_CHOICE,
                max_attempts=1,
                metadata={'explanation': q.get('explanation', '')}
            )
            order_index += 1
        
        # Summary step
        if content.get('summary'):
            LessonStep.objects.create(
                lesson=lesson,
                order_index=order_index,
                step_type=LessonStep.StepType.SUMMARY,
                teacher_script=content['summary'],
                answer_type=LessonStep.AnswerType.NONE,
            )
        
        # Store image suggestions in lesson metadata
        if content.get('image_suggestions'):
            lesson.metadata = lesson.metadata or {}
            lesson.metadata['image_suggestions'] = content['image_suggestions']
            lesson.save(update_fields=['metadata'])
    
    def _generate_lesson_media(self, course: Course):
        """Generate images for lessons using DALL-E."""
        
        if not self.dalle:
            print("    Skipping media generation (no OpenAI key)")
            return
        
        lessons = Lesson.objects.filter(unit__course=course)
        
        for lesson in lessons:
            suggestions = (lesson.metadata or {}).get('image_suggestions', [])
            
            for suggestion in suggestions[:2]:  # Max 2 images per lesson
                try:
                    print(f"      Generating image: {suggestion['title']}")
                    
                    # Generate image
                    response = self.dalle.images.generate(
                        model="dall-e-3",
                        prompt=suggestion['prompt'] + ". Style: educational illustration, clear and simple, suitable for secondary school students.",
                        size="1024x1024",
                        quality="standard",
                        n=1,
                    )
                    
                    image_url = response.data[0].url
                    
                    # Download image
                    import requests
                    image_response = requests.get(image_url)
                    image_response.raise_for_status()
                    
                    # Create MediaAsset
                    safe_title = "".join(c if c.isalnum() else "_" for c in suggestion['title'])
                    filename = f"{lesson.id}_{safe_title}.png"
                    
                    media_asset = MediaAsset.objects.create(
                        institution=self.institution,
                        title=suggestion['title'],
                        asset_type='image',
                        alt_text=suggestion.get('purpose', suggestion['title']),
                        caption=suggestion.get('purpose', ''),
                        tags=f"{course.title.lower()}, {lesson.unit.title.lower()}, ai-generated",
                    )
                    
                    media_asset.file.save(
                        filename,
                        ContentFile(image_response.content),
                        save=True
                    )
                    
                    # Link to first lesson step
                    first_step = lesson.steps.first()
                    if first_step:
                        StepMedia.objects.create(
                            lesson_step=first_step,
                            media_asset=media_asset,
                            placement='top',
                            order_index=0,
                        )
                    
                    # Rate limiting
                    time.sleep(2)
                    
                except Exception as e:
                    print(f"        Warning: Failed to generate image: {e}")


class Command(BaseCommand):
    """Django management command for curriculum generation."""
    
    help = 'Generate curriculum from a syllabus document'
    
    def add_arguments(self, parser):
        parser.add_argument('--file', type=str, required=True, help='Path to syllabus document')
        parser.add_argument('--subject', type=str, required=True, help='Subject name')
        parser.add_argument('--grade-level', type=str, default='S1-S5', help='Grade level')
        parser.add_argument('--institution', type=str, default='seychelles-secondary', help='Institution slug')
        parser.add_argument('--no-media', action='store_true', help='Skip media generation')
    
    def handle(self, *args, **options):
        from apps.accounts.models import Institution
        
        # Get API keys
        anthropic_key = os.environ.get('ANTHROPIC_API_KEY')
        openai_key = os.environ.get('OPENAI_API_KEY')
        
        if not anthropic_key:
            raise CommandError('ANTHROPIC_API_KEY not set')
        
        # Get institution
        try:
            institution = Institution.objects.get(slug=options['institution'])
        except Institution.DoesNotExist:
            raise CommandError(f"Institution '{options['institution']}' not found")
        
        # Generate curriculum
        generator = CurriculumGenerator(
            institution=institution,
            anthropic_key=anthropic_key,
            openai_key=openai_key,
        )
        
        course = generator.generate_from_document(
            document_path=options['file'],
            subject=options['subject'],
            grade_level=options['grade_level'],
            generate_media=not options['no_media'],
        )
        
        self.stdout.write(self.style.SUCCESS(f"✅ Created course: {course.title}"))
